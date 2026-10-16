"""Build the Excel workbook: Indonesia's travel services balance by partner country.

Main source = BPS reconstruction (one institution, one method, observed data). Mirror data (OECD TIS,
Eurostat, ABS), OECD-WTO BaTIS and destination-country receipts are shown only as comparisons.

Run:  python -I scripts/build_excel.py
Reads 03_extract/*.csv and 02_raw/fx/wb_PA.NUS.FCRF.json (via travel_data.py);
writes output/Neraca_Jasa_Perjalanan_Indonesia_per_Negara_v3.xlsx and 04_qa/reconciliation.md.

Every value column carries a status label (grey row under the header), defined in the 'Legenda' sheet:
GIVEN, GIVEN – MIRRORING, SELF-CALCULATED, SELF-CALCULATED – MIRRORING, SELF-ASSUMED.
"""
import sys
from pathlib import Path

sys.path.insert(0, str(Path(__file__).resolve().parent))

import pandas as pd  # noqa: E402
from openpyxl import Workbook  # noqa: E402
from openpyxl.chart import BarChart, Reference  # noqa: E402
from openpyxl.chart.label import DataLabelList  # noqa: E402
from openpyxl.styles import Alignment, Border, Font, PatternFill, Side  # noqa: E402
from openpyxl.utils import get_column_letter  # noqa: E402

from sources import INVENTORY, SRC  # noqa: E402
from travel_data import (EX, FX, ROOT, YEAR, batis_year, bps_credit, bps_debit,  # noqa: E402
                         destination_side, load_abs, load_batis, load_eurostat, load_oecd_tis, nama, pivot_cdb)

OUT = ROOT / "output" / "Neraca_Jasa_Perjalanan_Indonesia_per_Negara_v3.xlsx"
NAVY, ACCENT, GREY = "051C2C", "2251FF", "A2AAAD"
FONT = "Arial"
NUM = '#,##0.0;(#,##0.0);"–"'
NUM0 = '#,##0;(#,##0);"–"'
PCT = '0%;-0%;"–"'
LINK_CHECK = ROOT / "04_qa" / "link_check.csv"

H_KREDIT = "Kredit Indonesia\n(belanja warga negara ini di Indonesia)"
H_DEBIT = "Debit Indonesia\n(belanja penduduk Indonesia di negara ini)"
H_NERACA = "Neraca Indonesia\n(kredit − debit)"


# ---------------------------------------------------------------- styling helpers
def style_ws(ws):
    ws.sheet_view.showGridLines = False


def title(ws, text, sub=None):
    ws.cell(row=1, column=1, value=text).font = Font(name=FONT, size=14, bold=True, color=NAVY)
    if sub:
        ws.cell(row=2, column=1, value=sub).font = Font(name=FONT, size=9, italic=True, color="595959")


def bold(ws, row, col, text, size=11):
    ws.cell(row=row, column=col, value=text).font = Font(name=FONT, size=size, bold=True, color=NAVY)


def header(ws, row, col, labels, widths=None, height=42, prov=None, prov_height=40):
    """Writes a header row. `prov` = provenance labels for each column (list, None = blank) or one string merged
    across the value columns. Returns the first data row."""
    for i, lab in enumerate(labels):
        c = ws.cell(row=row, column=col + i, value=lab)
        c.font = Font(name=FONT, size=9, bold=True, color="FFFFFF")
        c.fill = PatternFill("solid", fgColor=NAVY)
        c.alignment = Alignment(horizontal="center", vertical="center", wrap_text=True)
        if widths:
            ws.column_dimensions[get_column_letter(col + i)].width = widths[i]
    ws.row_dimensions[row].height = height
    if prov is None:
        return row + 1
    pr = row + 1
    if isinstance(prov, str):
        prov = [None, prov] + [None] * (len(labels) - 2)
        ws.merge_cells(start_row=pr, start_column=col + 1, end_row=pr, end_column=col + len(labels) - 1)
    for i in range(len(labels)):
        c = ws.cell(row=pr, column=col + i, value=prov[i] if i < len(prov) else None)
        c.font = Font(name=FONT, size=8, italic=True, color="404040")
        c.fill = PatternFill("solid", fgColor="EDEDED")
        c.alignment = Alignment(horizontal="center", vertical="top", wrap_text=True)
    ws.row_dimensions[pr].height = prov_height
    return pr + 1


PROV_LEGEND = ("Status angka (baris abu-abu di bawah judul kolom): GIVEN = diambil apa adanya dari sumber · "
               "GIVEN – MIRRORING = laporan negara mitra, hanya sisi kredit/debit ditukar, angka tidak diubah · "
               "SELF-CALCULATED = kami hitung, rumusnya tertulis · SELF-CALCULATED – MIRRORING = kami hitung dari data "
               "negara mitra · SELF-ASSUMED = asumsi kami. Penjelasan lengkap: sheet 'Legenda'.")

# status tags: (tag, arti, contoh di workbook)
TAGS = [
    ("GIVEN", "Angka diambil apa adanya dari sumber resmi. Kami tidak menghitung, mengestimasi, atau mengubahnya.",
     "Kunjungan dan pengeluaran per kunjungan BPS; semua angka OECD-WTO BaTIS (F, N, B); total BI."),
    ("GIVEN – MIRRORING", "Angka dilaporkan negara mitra tentang Indonesia. Nilainya tidak diubah; hanya sisinya "
     "ditukar: impor jasa perjalanan negara mitra = kredit Indonesia, ekspor negara mitra = debit Indonesia.",
     "Kredit/debit di sheet Mirror – OECD TIS, Mirror – Eurostat (EUR), Mirror – ABS (AUD)."),
    ("SELF-CALCULATED", "Kami hitung sendiri dari angka GIVEN. Rumusnya ditulis di label kolom.",
     "Kredit BPS = kunjungan × pengeluaran per kunjungan; neraca = kredit − debit; konversi kurs ke USD."),
    ("SELF-CALCULATED – MIRRORING", "Kami hitung sendiri dari data negara mitra (mis. konversi kurs atau membalik "
     "tanda neraca mitra menjadi neraca Indonesia).", "Neraca Indonesia = neraca negara pelapor × (−1)."),
    ("SELF-ASSUMED", "Pilihan/asumsi kami karena sumber tidak menyediakan angkanya. Selalu disebut alasannya.",
     "Negara tanpa angka pengeluaran sendiri memakai rata-rata kelompoknya (BPS debit)."),
]

# OECD-WTO BaTIS codes, official labels from the 'label' columns of the BaTIS CSV (02_raw/batis/)
BATIS_CODES = [
    ("N", "Adjustment", "Reported", "Dilaporkan langsung oleh negara (untuk Indonesia: hanya total dunia)."),
    ("F", "Adjustment", "Adjusted and/or imputed", "Estimasi OECD-WTO untuk sisi Indonesia (diisi oleh OECD-WTO)."),
    ("B", "Adjustment", "Balanced", "Nilai seimbang buatan OECD-WTO: rata-rata tertimbang F dan laporan mitra."),
    ("ME", "Methodology type", "Model-based estimation",
     "Diisi OECD-WTO dengan model gravitasi (regresi PPML). Bukan model kami."),
    ("IE", "Methodology type", "Interpolation / extrapolation",
     "Diisi OECD-WTO dari deret waktu tahun lain (2023–2024 Indonesia sebagian besar IE)."),
    ("DE", "Methodology type", "Derivation", "Diturunkan OECD-WTO dari angka lain yang tersedia."),
    ("AG", "Methodology type", "Aggregation", "Penjumlahan OECD-WTO (mis. total dunia)."),
    ("I", "Observation status", "Value imputed by a receiving agency", "Nilai diisi OECD-WTO, bukan laporan negara."),
    ("A", "Observation status", "Normal value", "Nilai normal (laporan negara)."),
]
METHOD_LABEL = {c: f"{c} = {o}" for c, k, o, _ in BATIS_CODES if k == "Methodology type"}


def code_table(ws, row, rows, title_text, spans=((1, 1), (2, 2), (3, 4), (5, 9))):
    """Legend table with merged cells: code | kind | official label | plain meaning. Returns next free row."""
    bold(ws, row, 1, title_text, 10)
    labels = ["Kode", "Jenis kode", "Label resmi OECD-WTO", "Arti sederhana"]
    for (c1, c2), lab in zip(spans, labels):
        c = ws.cell(row=row + 1, column=c1, value=lab)
        c.font = Font(name=FONT, size=9, bold=True, color="FFFFFF")
        c.fill = PatternFill("solid", fgColor=NAVY)
        if c2 > c1:
            ws.merge_cells(start_row=row + 1, start_column=c1, end_row=row + 1, end_column=c2)
    thin = Side(style="hair", color=GREY)
    for i, vals in enumerate(rows):
        rr = row + 2 + i
        for (c1, c2), v in zip(spans, vals):
            c = ws.cell(row=rr, column=c1, value=v)
            c.font = Font(name=FONT, size=9, bold=(c1 == 1))
            c.border = Border(bottom=thin)
            c.alignment = Alignment(wrap_text=True, vertical="top")
            if c2 > c1:
                ws.merge_cells(start_row=rr, start_column=c1, end_row=rr, end_column=c2)
        ws.row_dimensions[rr].height = 24
    return row + 2 + len(rows)


def prov_legend(ws, row=3):
    c = ws.cell(row=row, column=1, value=PROV_LEGEND)
    c.font = Font(name=FONT, size=9, bold=True, color="404040")


def write_rows(ws, row, col, rows, fmts):
    thin = Side(style="hair", color=GREY)
    for r_i, r in enumerate(rows):
        for c_i, v in enumerate(r):
            if isinstance(v, float) and pd.isna(v):
                v = None
            c = ws.cell(row=row + r_i, column=col + c_i, value=v)
            c.font = Font(name=FONT, size=10)
            c.border = Border(bottom=thin)
            if fmts[c_i]:
                c.number_format = fmts[c_i]
    return row + len(rows)


def link(cell, url, text):
    cell.value = text
    cell.hyperlink = url
    cell.font = Font(name=FONT, size=9, color=ACCENT, underline="single")


def link_status():
    """url -> (verdict, date) from 04_qa/link_check.csv (written by scripts/check_links.py)."""
    if not LINK_CHECK.exists():
        return {}
    d = pd.read_csv(LINK_CHECK)
    return {r.url: (r.verdict, r.checked_on) for r in d.itertuples()}


def src_block(ws, row, keys, col=1):
    """Source block: per source a label row, then 'Halaman data' (browse page / download menu) and 'Unduh langsung'
    (direct file download) rows. Returns the next free row."""
    shown = set()
    for key in keys:
        s = SRC[key]
        c = ws.cell(row=row, column=col, value=f"Sumber: {s['label']}")
        c.font = Font(name=FONT, size=9, bold=True, color="404040")
        row += 1
        for kind, items in (("Halaman data", s["browse"]), ("Unduh langsung", s["unduh"])):
            for i, (text, url) in enumerate(items):
                ws.cell(row=row, column=col, value=kind if i == 0 else None).font = Font(name=FONT, size=9,
                                                                                         color="595959")
                link(ws.cell(row=row, column=col + 1), url, text)
                row += 1
        if s.get("catatan") and s["catatan"] not in shown:
            note(ws, row, s["catatan"], col=col + 1)
            shown.add(s["catatan"])
            row += 1
    return row


def note(ws, row, text, col=1):
    c = ws.cell(row=row, column=col, value=text)
    c.font = Font(name=FONT, size=9, italic=True, color="595959")


def bar_chart(ws, anchor, cat_col, val_col, r1, r2, chart_title):
    """Horizontal bar chart, one colour (no colour coding), largest bar on top."""
    ch = BarChart()
    ch.type = "bar"
    ch.title = chart_title
    ch.title.overlay = False
    ch.legend = None
    ch.add_data(Reference(ws, min_col=val_col, min_row=r1, max_row=r2), titles_from_data=False)
    ch.set_categories(Reference(ws, min_col=cat_col, min_row=r1, max_row=r2))
    s = ch.series[0]
    s.invertIfNegative = False
    s.graphicalProperties.solidFill = NAVY
    s.graphicalProperties.line.solidFill = NAVY
    ch.dataLabels = DataLabelList()
    ch.dataLabels.showVal = True
    for attr in ("showSerName", "showCatName", "showLegendKey"):
        setattr(ch.dataLabels, attr, False)
    ch.x_axis.scaling.orientation = "maxMin"
    ch.x_axis.tickLblPos = "low"
    ch.x_axis.delete = False
    ch.y_axis.delete = True
    ch.y_axis.majorGridlines = None
    ch.gapWidth = 60
    ch.width, ch.height = 14, 7
    ws.add_chart(ch, anchor)


# ---------------------------------------------------------------- top-5 logic
def top5(df, col, ascending=False):
    s = df[col].dropna()
    s = s[s != 0]
    return s.sort_values(ascending=ascending).head(5)


def surplus5(df):
    t = top5(df, "neraca")
    return t[t > 0]


def add_charts(ws, df, col, row, source_label, flows=("kredit", "debit", "neraca")):
    """Small top-5 tables (chart data) + one-colour charts next to them."""
    bold(ws, row, col, f"Data grafik: Top 5, {YEAR}, USD juta (salinan dari tabel utama)", 10)
    r = row + 1
    chart_row = row
    for flow in flows:
        if flow == "neraca":
            pos = surplus5(df)
            neg = top5(df, "neraca", ascending=True)
            neg = neg[neg < 0].sort_values(ascending=False)
            t = pd.concat([pos, neg])
        else:
            t = top5(df, flow)
        if t.empty:
            continue
        lab = {"kredit": "Kredit", "debit": "Debit", "neraca": "Neraca"}[flow]
        header(ws, r, col, ["Negara", lab], height=18)
        first = r + 1
        r = write_rows(ws, first, col, [[df.loc[i, "nama"], float(v)] for i, v in t.items()], [None, NUM0])
        lead = df.loc[t.index[0], "nama"]
        if flow == "kredit":
            line1 = f"Wisatawan dari {lead} paling banyak belanja di Indonesia"
            line2 = "Kredit Indonesia"
        elif flow == "debit":
            line1 = f"Penduduk Indonesia paling banyak belanja di {lead}"
            line2 = "Debit Indonesia"
        else:
            worst = df.loc[t.index[-1], "nama"] if t.iloc[-1] < 0 else None
            line1 = f"Surplus terbesar dengan {lead}" + (f", defisit terbesar dengan {worst}" if worst else "")
            line2 = "Neraca Indonesia = kredit − debit; batang ke kiri = defisit"
        bar_chart(ws, f"{get_column_letter(col + 3)}{chart_row}", col, col + 1, first, r - 1,
                  f"{line1}\n{line2} · Top 5 · USD juta · {YEAR} · {source_label}")
        chart_row += 15
        r += 2
    ws.column_dimensions[get_column_letter(col)].width = 20
    ws.column_dimensions[get_column_letter(col + 1)].width = 10


# ================================================================ BUILD
def main():
    wb = Workbook()
    wb.remove(wb.active)

    # ---------- data (USD m, Indonesia's perspective, YEAR)
    tis = load_oecd_tis()
    tis_y = pivot_cdb(tis[tis.year == YEAR], "reporter")
    tis_y.columns = ["kredit", "debit", "neraca"]
    tis_y["nama"] = [nama(i) for i in tis_y.index]

    eu = load_eurostat()
    ey = eu[eu.year == YEAR]
    usd = ey.pivot_table(index="reporter", columns="flow_idn", values="value_usd", aggfunc="sum")
    eur = ey.pivot_table(index="reporter", columns="flow_idn", values="value", aggfunc="sum")
    eu_y = pd.DataFrame({"kredit": usd.get("credit"), "debit": usd.get("debit"), "neraca": usd.get("balance"),
                         "kredit_eur": eur.get("credit"), "debit_eur": eur.get("debit"),
                         "neraca_eur": eur.get("balance")}).dropna(how="all", subset=["kredit", "debit"])
    eu_y["nama"] = [nama(i) for i in eu_y.index]

    batis = load_batis()
    bt = batis_year(batis, YEAR)
    bt_std = pd.DataFrame({"kredit": bt.kredit_B, "debit": bt.debit_B, "neraca": bt.neraca_B, "nama": bt.nama})

    cr, cr_meta = bps_credit(YEAR)
    cr["kredit"] = cr.kunjungan * cr.pengeluaran_per_kunjungan / 1e6
    de, de_resid, de_meta = bps_debit(YEAR)
    de["perjalanan"] = de_meta["total_perjalanan"] * de.pangsa_pct / 100
    de["debit"] = de.perjalanan * de.pengeluaran_per_perjalanan / 1e6
    bps = pd.DataFrame(index=sorted(set(cr.index) | set(de.index)))
    bps["nama"] = [nama(i) for i in bps.index]
    bps["kredit"] = cr.kredit.reindex(bps.index)
    bps["debit"] = de.debit.reindex(bps.index)
    bps["neraca"] = bps.kredit - bps.debit
    bps["basis_debit"] = de.sumber_pengeluaran.reindex(bps.index)
    bps["kelompok"] = de.kelompok_pengeluaran.reindex(bps.index)
    bc = float(cr_meta["total_kunjungan"] * cr_meta["rata_rata"] / 1e6)
    bd = float(de.debit.sum() + (de_meta["total_perjalanan"] * de_resid.pangsa_pct / 100 *
                                 de_resid.pengeluaran_per_perjalanan / 1e6).sum())

    ds = destination_side(YEAR)
    ds_std = pd.DataFrame({"kredit": float("nan"), "debit": ds.debit_usd, "neraca": float("nan"), "nama": ds.nama})

    bi = pd.read_csv(EX / "bi_travel_totals.csv")
    bi = bi[(bi["item"] == "travel") & bi.quarter.isna()].pivot_table(index="year", columns="flow_idn", values="value")

    # ================================================================ 0. Legenda
    ws = wb.create_sheet("Legenda")
    style_ws(ws)
    title(ws, "Legenda: dari mana setiap angka berasal",
          "Setiap kolom angka di workbook ini punya label status di baris abu-abu tepat di bawah judul kolom.")
    header(ws, 4, 1, ["Label status", "Artinya", "Contoh di workbook ini"], [30, 70, 70], height=20)
    end = write_rows(ws, 5, 1, [list(t) for t in TAGS], [None] * 3)
    for rr in range(5, end):
        ws.cell(row=rr, column=1).font = Font(name=FONT, size=10, bold=True, color=NAVY)
        for cc in (2, 3):
            ws.cell(row=rr, column=cc).alignment = Alignment(wrap_text=True, vertical="top")
        ws.row_dimensions[rr].height = 42
    r = end + 1
    bold(ws, r, 1, "Ringkasan per sheet", 10)
    header(ws, r + 1, 1, ["Sheet", "Status angka", "Yang kami hitung sendiri"], height=20)
    per_sheet = [
        ["Neraca per Negara / BPS – Perhitungan", "Input GIVEN (BPS); hasil SELF-CALCULATED",
         "Kredit = kunjungan × pengeluaran per kunjungan; debit = total perjalanan × pangsa negara × pengeluaran per "
         "perjalanan; negara tanpa angka pengeluaran sendiri = SELF-ASSUMED (rata-rata kelompok)."],
        ["Mirror – OECD TIS / Eurostat / ABS", "GIVEN – MIRRORING (nilai lokal); SELF-CALCULATED – MIRRORING",
         "Neraca = neraca negara pelapor × (−1) atau kredit − debit; konversi EUR/AUD ke USD dengan kurs Bank Dunia."],
        ["BaTIS Data / BaTIS – Cara Model", "GIVEN: semua kredit dan debit diambil langsung dari OECD-WTO",
         "Tidak ada estimasi kami. Hanya neraca = kredit B − debit B (BaTIS tidak menerbitkan neraca)."],
        ["Data Negara Tujuan", "Nilai lokal GIVEN; USD SELF-CALCULATED; haji SELF-CALCULATED + SELF-ASSUMED",
         "Konversi ke USD; haji = kuota × (BPIH − tiket); biaya tiket 2025 dipakai untuk tahun lain."],
        ["Tren / Rekonsiliasi / Perbandingan Sumber", "Disalin dari sheet lain; rasio dan jumlah SELF-CALCULATED",
         "Neraca per tahun, jumlah semua mitra, rasio terhadap BI."],
    ]
    end = write_rows(ws, r + 2, 1, per_sheet, [None] * 3)
    for rr in range(r + 2, end):
        for cc in (2, 3):
            ws.cell(row=rr, column=cc).alignment = Alignment(wrap_text=True, vertical="top")
        ws.row_dimensions[rr].height = 42
    code_table(ws, end + 1, [list(x) for x in BATIS_CODES],
               "Kode OECD-WTO BaTIS (ME, IE, DE, AG, dan lain-lain): label resmi dari file BaTIS",
               spans=((1, 1), (2, 2), (3, 3), (4, 4)))
    ws.column_dimensions["D"].width = 70

    # ================================================================ 1. Neraca per Negara (BPS)
    ws = wb.create_sheet("Neraca per Negara")
    style_ws(ws)
    title(ws, f"Neraca jasa perjalanan Indonesia per negara, {YEAR} (USD juta) — sumber utama: BPS",
          "Kredit = belanja wisatawan mancanegara di Indonesia (ekspor jasa). Debit = belanja penduduk Indonesia di "
          "luar negeri (impor jasa). Neraca = kredit − debit; angka dalam kurung = defisit.")
    bold(ws, 4, 1, "Kenapa BPS yang dipakai, bukan dicampur:", 10)
    why = [
        "1. Data observasi, bukan model: jumlah kunjungan/perjalanan (Imigrasi, data seluler) × survei pengeluaran "
        "(Passenger Exit Survey dan Survei Digital Wisnas).",
        "2. Kredit dan debit dari satu lembaga dengan satu metode, mencakup pasar utama kedua arah. Mirror hanya ada "
        "untuk 33 negara pelapor (tanpa ASEAN, Tiongkok, Arab Saudi); BaTIS hasil model (debit ke Arab Saudi cuma "
        f"USD {bt.loc['SAU', 'debit_B']:,.0f} juta).",
        f"3. Bahan yang sama dengan neraca perjalanan BI (PES). Catatan level: total BPS {bc / bi.loc[YEAR, 'credit']:.0%}"
        f" (kredit) dan {bd / bi.loc[YEAR, 'debit']:.0%} (debit) dari angka BI, jadi pakai untuk urutan & porsi negara.",
    ]
    for i, t in enumerate(why):
        note(ws, 5 + i, t)
    note(ws, 8, "PES = Passenger Exit Survey, survei untuk menghitung neraca perjalanan (Kemenpar). Link di bawah tabel.")
    prov_legend(ws)
    first = header(ws, 10, 1, ["Negara", H_KREDIT, H_DEBIT, H_NERACA, "Catatan"], [22, 16, 16, 14, 46], height=48,
                   prov=[None, "SELF-CALCULATED: kunjungan × pengeluaran per kunjungan (input GIVEN BPS)",
                         "SELF-CALCULATED: perjalanan × pangsa tujuan × pengeluaran per perjalanan (input GIVEN BPS)",
                         "SELF-CALCULATED: kredit − debit", "Catatan kami"], prov_height=52)
    main_tbl = bps.copy()
    main_tbl["_sort"] = main_tbl.neraca.fillna(-1e12)
    main_tbl = main_tbl.sort_values("_sort", ascending=False)

    def catatan(r):
        if pd.isna(r.kredit):
            return "Kredit tidak dirinci BPS (masuk 'negara lainnya')"
        if pd.isna(r.debit):
            return "Debit tidak dirinci BPS (masuk tujuan 'lainnya')"
        if r.basis_debit == "rata-rata kelompok":
            return f"Debit memakai rata-rata belanja kelompok '{r.kelompok}'"
        return ""
    rows = [[r.nama, r.kredit, r.debit, r.neraca, catatan(r)] for _, r in main_tbl.iterrows()]
    end = write_rows(ws, first, 1, rows, [None, NUM, NUM, NUM, None])
    ws.freeze_panes = f"A{first}"
    note(ws, end + 1, "Rincian perhitungan ada di sheet 'BPS – Perhitungan'.")
    src_block(ws, end + 3, ["BPS_KREDIT_KUNJUNGAN", "BPS_KREDIT_PENGELUARAN", "BPS_PERJALANAN", "BPS_DEBIT_PANGSA",
                          "BPS_DEBIT_PENGELUARAN", "PES"])
    add_charts(ws, bps, 7, 10, "Sumber: BPS")

    # ================================================================ 2. Perbandingan Sumber
    ws = wb.create_sheet("Perbandingan Sumber")
    style_ws(ws)
    title(ws, f"Apakah Top 5 negara sama di setiap sumber? ({YEAR}, USD juta)",
          "BPS = sumber utama. Sumber lain hanya pembanding. Angka dalam kurung = nilai, disalin dari sheet masing-"
          "masing sumber (status GIVEN/DIHITUNG ada di sheet itu).")
    prov_legend(ws)
    leg = ws.cell(row=4, column=1, value="Legenda")
    leg.font = Font(name=FONT, size=9, bold=True)
    c = ws.cell(row=4, column=2, value="Nama negara")
    c.fill = PatternFill("solid", fgColor=ACCENT)
    c.font = Font(name=FONT, size=9, bold=True, color="FFFFFF")
    note(ws, 4, "← artinya negara ini masuk Top 5 di minimal 3 dari 5 sumber (hasilnya konsisten). Tanpa warna = hanya muncul "
                "di 1–2 sumber.", col=3)
    sources = {"BPS (utama)": bps, "Mirror OECD TIS": tis_y, "Mirror Eurostat": eu_y, "Model BaTIS": bt_std,
               "Data negara tujuan": ds_std}
    consensus = {}
    for flow in ("kredit", "debit", "neraca"):
        cnt = {}
        for d in sources.values():
            idx = surplus5(d).index if flow == "neraca" else top5(d, flow).index
            for i in idx:
                cnt[i] = cnt.get(i, 0) + 1
        consensus[flow] = {i for i, n_ in cnt.items() if n_ >= 3}
    r = 6
    for flow, lab in (("kredit", "Kredit: belanja wisatawan dari negara ini di Indonesia"),
                      ("debit", "Debit: belanja penduduk Indonesia di negara ini"),
                      ("neraca", "Surplus terbesar (kredit − debit)")):
        bold(ws, r, 1, lab, 10)
        names = [n for n in sources if not (flow != "debit" and n == "Data negara tujuan")]
        header(ws, r + 1, 1, ["Peringkat"] + names, [10] + [24] * len(names), height=30)
        lists = {n: list((surplus5(sources[n]) if flow == "neraca" else top5(sources[n], flow)).index) for n in names}
        for k in range(5):
            ws.cell(row=r + 2 + k, column=1, value=k + 1).font = Font(name=FONT, size=10)
            for j, n in enumerate(names):
                if k < len(lists[n]):
                    iso = lists[n][k]
                    d = sources[n]
                    cc = ws.cell(row=r + 2 + k, column=2 + j, value=f"{d.loc[iso, 'nama']} ({d.loc[iso, flow]:,.0f})")
                    hit = iso in consensus[flow]
                    cc.font = Font(name=FONT, size=10, bold=hit, color="FFFFFF" if hit else "000000")
                    if hit:
                        cc.fill = PatternFill("solid", fgColor=ACCENT)
        rr = r + 7
        ws.cell(row=rr, column=1, value="Sama dengan BPS").font = Font(name=FONT, size=9, italic=True)
        ref = set(lists["BPS (utama)"])
        for j, n in enumerate(names):
            ws.cell(row=rr, column=2 + j, value=f"{len(set(lists[n]) & ref)} dari 5").font = Font(
                name=FONT, size=9, italic=True)
        r = rr + 2
    note(ws, r, "Kenapa bisa beda: mirror hanya berisi negara pelapor (OECD TIS tanpa Tiongkok, Malaysia, Singapura, "
                "Arab Saudi, Korea, Thailand; Eurostat hanya Eropa). BaTIS mencakup semua negara tapi hasil model.")
    # cross-checks
    r += 2
    bold(ws, r, 1, "Uji silang debit BPS dengan data resmi negara tujuan", 10)
    first = header(ws, r + 1, 1, ["Negara", "Debit BPS", "Data negara tujuan", "BPS / negara tujuan",
                                  "Catatan negara tujuan"], height=30,
                   prov=[None, "Dari sheet Neraca per Negara (SELF-CALCULATED)", "Dari sheet Data Negara Tujuan",
                         "SELF-CALCULATED: debit BPS ÷ data negara tujuan", None])
    rows = []
    for iso, x in ds.iterrows():
        b = bps.debit.get(iso)
        rows.append([x.nama, b, x.debit_usd, b / x.debit_usd if b else None, f"{x.sumber}: {x.catatan or x.periode}"])
    r = write_rows(ws, first, 1, rows, [None, NUM0, NUM0, PCT, None])
    r += 1
    bold(ws, r, 1, "Uji silang kredit BPS dengan laporan negara mitra (mirror OECD TIS)", 10)
    first = header(ws, r + 1, 1, ["Negara", "Kredit BPS", "Kredit menurut negara mitra", "BPS / mitra", ""],
                   height=30, prov=[None, "Dari sheet Neraca per Negara (SELF-CALCULATED)",
                                    "Dari sheet Mirror – OECD TIS (GIVEN – MIRRORING)",
                                    "SELF-CALCULATED: kredit BPS ÷ kredit mitra", None])
    both = [i for i in bps.kredit.dropna().sort_values(ascending=False).index if i in tis_y.index][:8]
    rows = [[bps.loc[i, "nama"], bps.loc[i, "kredit"], tis_y.loc[i, "kredit"],
             bps.loc[i, "kredit"] / tis_y.loc[i, "kredit"], ""] for i in both]
    r = write_rows(ws, first, 1, rows, [None, NUM0, NUM0, PCT, None])
    note(ws, r + 1, "Australia melaporkan belanja warganya di Indonesia jauh lebih besar dari survei BPS; selisih "
                    "antarnegara seperti ini (asimetri) umum di statistik jasa.")

    # ================================================================ 3. BPS – Perhitungan
    ws = wb.create_sheet("BPS – Perhitungan")
    style_ws(ws)
    title(ws, f"Perhitungan neraca dari data BPS, {YEAR}",
          "Kredit = kunjungan wisman × rata-rata pengeluaran per kunjungan. Debit = total perjalanan WNI × pangsa "
          "negara tujuan × rata-rata pengeluaran per perjalanan (tanpa tiket pesawat internasional).")
    prov_legend(ws)
    bold(ws, 4, 1, "Kredit: belanja wisatawan mancanegara di Indonesia")
    c0 = header(ws, 5, 1, ["Negara (kebangsaan)", "Kunjungan (orang)", "Pengeluaran per kunjungan (USD)",
                           "Kredit (USD juta)"], [26, 14, 16, 14],
                prov=[None, "GIVEN (BPS var 1470, jumlah Jan–Des)", "GIVEN (BPS var 272)",
                      "SELF-CALCULATED: B × C ÷ 1.000.000"])
    r = c0
    for _, row in cr.sort_values("kredit", ascending=False).iterrows():
        write_rows(ws, r, 1, [[row.nama, row.kunjungan, row.pengeluaran_per_kunjungan, None]], [None, NUM0, NUM, NUM])
        ws.cell(row=r, column=4, value=f"=B{r}*C{r}/1000000").number_format = NUM
        r += 1
    write_rows(ws, r, 1, [["Negara lainnya (sisa)", f"=B{r + 1}-SUM(B{c0}:B{r - 1})", None, None]],
               [None, NUM0, NUM, NUM])
    ws.cell(row=r, column=4, value=f"=D{r + 1}-SUM(D{c0}:D{r - 1})").number_format = NUM
    write_rows(ws, r + 1, 1, [["Total semua kunjungan", cr_meta["total_kunjungan"], cr_meta["rata_rata"], None]],
               [None, NUM0, NUM, NUM])
    ws.cell(row=r + 1, column=4, value=f"=B{r + 1}*C{r + 1}/1000000").number_format = NUM
    for cc in (1, 4):
        ws.cell(row=r + 1, column=cc).font = Font(name=FONT, size=10, bold=True)
    note(ws, r + 2, "Baris total: kunjungan dan rata-rata (USD 1.391,85) GIVEN dari BPS; kredit SELF-CALCULATED "
                    "(B × C). Baris 'negara lainnya': SELF-CALCULATED = total dikurangi negara yang dirinci.")
    r = src_block(ws, r + 4, ["BPS_KREDIT_KUNJUNGAN", "BPS_KREDIT_PENGELUARAN"]) + 1
    bold(ws, r, 1, "Debit: belanja penduduk Indonesia di luar negeri")
    ws.cell(row=r + 1, column=1, value="Total perjalanan WNI ke luar negeri (GIVEN, BPS var 2195)").font = Font(
        name=FONT, size=10)
    ws.cell(row=r + 1, column=2, value=de_meta["total_perjalanan"]).number_format = NUM0
    tot = f"$B${r + 1}"
    d0 = header(ws, r + 2, 1, ["Negara tujuan", "Pangsa perjalanan (%)", "Pengeluaran per perjalanan (USD)",
                               "Debit (USD juta)", "Basis pengeluaran"],
                prov=[None, "GIVEN (BPS SWN 2025, Lampiran 4)",
                      "GIVEN (BPS SWN 2024, Lampiran 20) — kecuali baris 'rata-rata kelompok' = SELF-ASSUMED",
                      "SELF-CALCULATED: total perjalanan × B ÷ 100 × C ÷ 1.000.000",
                      "SELF-ASSUMED: negara tanpa angka sendiri dipetakan ke kelompoknya"], prov_height=52)
    ws.column_dimensions["E"].width = 30
    r = d0
    for _, row in de.sort_values("debit", ascending=False).iterrows():
        basis = "angka negara ini" if row.sumber_pengeluaran == "rata-rata negara" else \
            f"rata-rata kelompok '{row.kelompok_pengeluaran}'"
        write_rows(ws, r, 1, [[row.nama, row.pangsa_pct, row.pengeluaran_per_perjalanan, None, basis]],
                   [None, "0.00", NUM, NUM, None])
        ws.cell(row=r, column=4, value=f"={tot}*B{r}/100*C{r}/1000000").number_format = NUM
        r += 1
    for _, row in de_resid.iterrows():
        write_rows(ws, r, 1, [[f"{row.label} (sisa)", row.pangsa_pct, row.pengeluaran_per_perjalanan, None,
                               f"rata-rata kelompok '{row.kelompok_pengeluaran}'"]], [None, "0.00", NUM, NUM, None])
        ws.cell(row=r, column=4, value=f"={tot}*B{r}/100*C{r}/1000000").number_format = NUM
        r += 1
    ws.cell(row=r, column=1, value="Total").font = Font(name=FONT, size=10, bold=True)
    ws.cell(row=r, column=4, value=f"=SUM(D{d0}:D{r - 1})").number_format = NUM
    ws.cell(row=r, column=4).font = Font(name=FONT, size=10, bold=True)
    note(ws, r + 1, "BPS hanya menerbitkan pengeluaran per perjalanan untuk 16 negara/kelompok; negara lain memakai "
                    "rata-rata kelompoknya. Paket wisata (mis. umrah) bisa ikut memuat tiket pesawat.")
    src_block(ws, r + 3, ["BPS_PERJALANAN", "BPS_DEBIT_PANGSA", "BPS_DEBIT_PENGELUARAN"])

    # ================================================================ 4. Mirror – OECD TIS
    ws = wb.create_sheet("Mirror – OECD TIS")
    style_ws(ws)
    title(ws, f"Pembanding: laporan negara mitra ke OECD, {YEAR} (USD juta)",
          "Angka dilaporkan negara pelapor dari sisi mereka, lalu dibalik ke sisi Indonesia: impor jasa perjalanan "
          "negara pelapor (belanja warganya di Indonesia) = kredit Indonesia, dan sebaliknya.")
    prov_legend(ws)
    d = tis_y.sort_values("neraca", ascending=False)
    first = header(ws, 4, 1, ["Negara pelapor", H_KREDIT, H_DEBIT, H_NERACA], [22, 16, 16, 14], height=48,
                   prov=[None, "GIVEN – MIRRORING: = debit (impor) negara pelapor, nilai tidak diubah",
                         "GIVEN – MIRRORING: = kredit (ekspor) negara pelapor, nilai tidak diubah",
                         "SELF-CALCULATED – MIRRORING: neraca negara pelapor × (−1)"], prov_height=52)
    end = write_rows(ws, first, 1, [[r_.nama, r_.kredit, r_.debit, r_.neraca] for _, r_ in d.iterrows()],
                     [None, NUM, NUM, NUM])
    src_block(ws, end + 2, ["OECD_TIS"])
    add_charts(ws, d, 7, 4, "Sumber: OECD TIS")

    # ================================================================ 5. Mirror – Eurostat
    ws = wb.create_sheet("Mirror – Eurostat")
    style_ws(ws)
    title(ws, f"Pembanding: laporan negara Eropa ke Eurostat, {YEAR}",
          "Sudah dibalik ke sisi Indonesia. EUR juta sesuai sumber; USD juta memakai kurs rata-rata Bank Dunia.")
    prov_legend(ws)
    d = eu_y.sort_values("neraca", ascending=False)
    conv = f"SELF-CALCULATED: EUR ÷ {FX[('EMU', YEAR)]:.4f} (kurs EUR per USD {YEAR}, Bank Dunia)"
    first = header(ws, 4, 1, ["Negara pelapor", "Kredit Indonesia (EUR juta)", "Debit Indonesia (EUR juta)",
                              "Neraca Indonesia (EUR juta)", H_KREDIT + " USD juta", H_DEBIT + " USD juta",
                              H_NERACA + " USD juta"], [20, 14, 14, 14, 16, 16, 14], height=60,
                   prov=[None, "GIVEN – MIRRORING: = debit negara pelapor", "GIVEN – MIRRORING: = kredit negara pelapor",
                         "SELF-CALCULATED – MIRRORING: neraca pelapor × (−1)", conv, conv, conv], prov_height=64)
    end = write_rows(ws, first, 1, [[r_.nama, r_.kredit_eur, r_.debit_eur, r_.neraca_eur, r_.kredit, r_.debit,
                                     r_.neraca] for _, r_ in d.iterrows()], [None] + [NUM] * 6)
    src_block(ws, end + 2, ["EUROSTAT", "WB_KURS"])
    add_charts(ws, d, 10, 4, "Sumber: Eurostat")

    # ================================================================ 6. Mirror – ABS
    ws = wb.create_sheet("Mirror – ABS")
    style_ws(ws)
    title(ws, "Pembanding: laporan Australia (ABS) tentang perjalanan dengan Indonesia",
          "Satu negara mitra, jadi tidak ada Top 5. Mata uang tabel ABS '$m' = juta dolar Australia (AUD).")
    prov_legend(ws)
    bold(ws, 4, 1, "Dari mana kredit dan debit Indonesia diambil", 10)
    header(ws, 5, 1, ["Sisi Indonesia", "Tabel ABS", "Arti", "Unduh langsung", "Halaman data"], [26, 34, 52, 14, 16],
           height=20)
    rows = [["Kredit Indonesia", "Tabel 8 'Debits' (impor jasa perjalanan Australia)",
             "Belanja penduduk Australia saat berada di Indonesia", None],
            ["Debit Indonesia", "Tabel 7 'Credits' (ekspor jasa perjalanan Australia)",
             "Belanja penduduk Indonesia saat berada di Australia (termasuk biaya kuliah)", None]]
    write_rows(ws, 6, 1, rows, [None] * 4)
    abs_dl = dict(SRC["ABS"]["unduh"])
    for rr, tab in ((6, "Tabel 8 (xlsx)"), (7, "Tabel 7 (xlsx)")):
        link(ws.cell(row=rr, column=4), abs_dl[tab], tab)
        link(ws.cell(row=rr, column=5), SRC["ABS"]["browse"][0][1], "Halaman rilis ABS")
    t = tis[(tis.reporter == "AUS")]
    note(ws, 8, "Bukti AUD: nilai ABS sama persis dengan nilai Australia dalam mata uang nasional di OECD (mis. 2024: "
                "10.223 dan 2.317), dan kurs tersirat ke USD 2023 = 1,5052 = kurs AUD Bank Dunia.")
    a = load_abs()
    a_sd = a[(a["item"] == "SD") & (a.year >= 2010)]
    p = a_sd.pivot_table(index="year", columns="flow_idn", values="value", aggfunc="first")
    pu = a_sd.pivot_table(index="year", columns="flow_idn", values="value_usd", aggfunc="first")
    conv = "SELF-CALCULATED: AUD ÷ kurs AUD per USD tahun itu (Bank Dunia)"
    a0 = header(ws, 10, 1, ["Tahun", "Kredit Indonesia (AUD juta)", "Debit Indonesia (AUD juta)",
                            "Neraca Indonesia (AUD juta)", "Kredit Indonesia (USD juta)", "Debit Indonesia (USD juta)",
                            "Neraca Indonesia (USD juta)"], [26, 14, 14, 14, 14, 14, 14], height=42,
                prov=[None, "GIVEN – MIRRORING: Tabel 8.1 (impor Australia), tanda minus dihapus",
                      "GIVEN – MIRRORING: Tabel 7.1 (ekspor Australia)", "SELF-CALCULATED: kredit − debit",
                      conv, conv, "SELF-CALCULATED: kredit − debit"], prov_height=52)
    rows = [[int(y), p.loc[y, "credit"], p.loc[y, "debit"], p.loc[y, "credit"] - p.loc[y, "debit"],
             pu.loc[y, "credit"], pu.loc[y, "debit"], pu.loc[y, "credit"] - pu.loc[y, "debit"]] for y in p.index]
    end = write_rows(ws, a0, 1, rows, [None] + [NUM0] * 6)
    yy = max(p.index)
    sub = a[(a.year == yy) & (a["item"].isin(["SDA", "SDB2", "ABS_OTHER_PERSONAL"]))]
    sp = sub.pivot_table(index="item", columns="flow_idn", values="value", aggfunc="first")
    lab = {"SDA": "Perjalanan bisnis", "SDB2": "Pendidikan (pribadi)",
           "ABS_OTHER_PERSONAL": "Pribadi lainnya (wisata, kesehatan, dll.)"}
    r0 = end + 2
    bold(ws, r0, 1, f"Rincian {yy} (AUD juta)", 10)
    f0 = header(ws, r0 + 1, 1, ["Jenis perjalanan", "Kredit Indonesia", "Debit Indonesia", "Neraca Indonesia"],
                height=20, prov=[None, "GIVEN – MIRRORING: Tabel 8.2 / 8.4 / 8.5", "GIVEN – MIRRORING: Tabel 7.2 / 7.4 / 7.5",
                                 "SELF-CALCULATED: kredit − debit"])
    f_end = write_rows(ws, f0, 1, [[lab[k], sp.loc[k, "credit"], sp.loc[k, "debit"],
                                    sp.loc[k, "credit"] - sp.loc[k, "debit"]] for k in lab if k in sp.index],
                       [None, NUM0, NUM0, NUM0])
    src_block(ws, f_end + 1, ["ABS", "WB_KURS"])
    neg_years = [int(y) for y in pu.index if pu.loc[y, "credit"] - pu.loc[y, "debit"] < 0]
    ch = BarChart()
    ch.type = "col"
    ch.title = ("Indonesia selalu surplus dengan Australia" + (
        f", kecuali {', '.join(map(str, neg_years))} (pandemi)" if neg_years else "") +
        "\nNeraca Indonesia, USD juta · Sumber: ABS")
    ch.title.overlay = False
    ch.legend = None
    ch.add_data(Reference(ws, min_col=7, min_row=a0, max_row=end - 1), titles_from_data=False)
    ch.set_categories(Reference(ws, min_col=1, min_row=a0, max_row=end - 1))
    ch.series[0].graphicalProperties.solidFill = NAVY
    ch.series[0].graphicalProperties.line.solidFill = NAVY
    ch.series[0].invertIfNegative = False
    ch.y_axis.majorGridlines = None
    ch.x_axis.delete = False
    ch.x_axis.tickLblPos = "low"
    ch.y_axis.delete = False
    ch.y_axis.numFmt = "#,##0"
    ch.width, ch.height = 18, 8
    ws.add_chart(ch, "I10")

    # ================================================================ 7. BaTIS Data
    ws = wb.create_sheet("BaTIS Data")
    style_ws(ws)
    title(ws, f"Pembanding: OECD-WTO BaTIS, {YEAR} (USD juta)",
          "SEMUA angka kredit dan debit di sheet ini GIVEN: diambil langsung dari OECD-WTO BaTIS, apa adanya. Tidak "
          "ada yang kami hitung atau estimasi sendiri; model gravitasi/PPML dijalankan OECD-WTO, bukan kami.")
    prov_legend(ws)
    note(ws, 4, "Satu-satunya operasi kami: kolom Neraca = kredit B − debit B (BaTIS tidak menerbitkan neraca). "
                "Tukar sisi ekspor/impor negara mitra di kolom N = mirroring, angkanya tidak diubah.")
    r_leg = code_table(ws, 6, [list(x) for x in BATIS_CODES],
                       "Legenda kode OECD-WTO (label resmi diambil dari file BaTIS; cara model: sheet 'BaTIS – Cara Model')")
    d = bt[(bt.kredit_B.fillna(0) + bt.debit_B.fillna(0)) > 0].sort_values("kredit_B", ascending=False)
    g = "GIVEN (OECD-WTO, kode {}) — diambil langsung"
    first = header(ws, r_leg + 1, 1, ["Negara", "Kredit Indonesia: estimasi OECD-WTO (F)",
                              "Kredit Indonesia: dilaporkan negara mitra (N)", "Kredit Indonesia: nilai seimbang (B)",
                              "Debit Indonesia: estimasi OECD-WTO (F)", "Debit Indonesia: dilaporkan negara mitra (N)",
                              "Debit Indonesia: nilai seimbang (B)", "Neraca Indonesia (B)", "Kode metode F",
                              "Arti kode metode"],
                   [22, 13, 13, 13, 13, 13, 13, 11, 10, 34], height=64,
                   prov=[None, g.format("F"), "GIVEN – MIRRORING (OECD-WTO, kode N): laporan mitra, angka tidak diubah",
                         g.format("B"), g.format("F"),
                         "GIVEN – MIRRORING (OECD-WTO, kode N): laporan mitra, angka tidak diubah", g.format("B"),
                         "SELF-CALCULATED: kredit B − debit B (pengurangan saja)", "GIVEN (OECD-WTO)",
                         "Legenda (lihat tabel di atas)"], prov_height=52)

    def arti(m):
        return " · ".join(METHOD_LABEL.get(x, x) for x in str(m).split("/")) if pd.notna(m) and m else None
    end = write_rows(ws, first, 1, [[r_.nama, r_.kredit_F, r_.kredit_mitra, r_.kredit_B, r_.debit_F, r_.debit_mitra,
                                     r_.debit_B, r_.neraca_B, r_.metode, arti(r_.metode)] for _, r_ in d.iterrows()],
                     [None] + [NUM] * 7 + [None, None])
    ws.freeze_panes = f"B{first}"
    note(ws, end + 1, "Kosong di kolom 'dilaporkan negara mitra' = negara itu tidak melapor ke OECD-WTO.")
    src_block(ws, end + 3, ["BATIS"])
    add_charts(ws, bt_std, 12, r_leg + 1, "Sumber: BaTIS, nilai seimbang (B)")

    # ================================================================ 8. BaTIS – Cara Model
    ws = wb.create_sheet("BaTIS – Cara Model")
    style_ws(ws)
    title(ws, "Bagaimana OECD-WTO membuat angka BaTIS (dan mengapa untuk Indonesia semuanya hasil model)",
          "Yang menghitung: OECD dan WTO. Kami hanya mengunduh hasilnya; tabel ini meringkas dokumen metodologi mereka.")
    prov_legend(ws)
    top = src_block(ws, 4, ["BATIS_METODE"])
    header(ws, top + 1, 1, ["Langkah", "Yang dilakukan OECD-WTO", "Dampak untuk angka Indonesia"], [22, 55, 70],
           height=20)
    steps = [
        ["A. Kumpulkan", "Mengumpulkan data bilateral yang dilaporkan semua negara",
         "Indonesia hanya melapor total dunia (2005–2024, 0 baris bilateral). Negara mitra (Australia, UE, AS, Jepang, "
         "dll.) melaporkan Indonesia."],
        ["B. Estimasi gravitasi", "Arus yang kosong diisi model gravitasi (regresi PPML) yang dijalankan OECD-WTO",
         "Variabel: PDB kedua negara, jarak, perbatasan, bahasa, hubungan kolonial, PDB per kapita, perdagangan jasa "
         "dengan dunia, tren waktu, perdagangan barang bilateral, efek tetap mitra, dan kedatangan/keberangkatan "
         "wisatawan (khusus perjalanan, data UN Tourism)."],
        ["C. Penskalaan", "Hasil estimasi diskalakan ke total dunia tiap negara",
         "Total sisi Indonesia = total dunia Indonesia (= angka BI). Pangsa per negara mitra tetap dari model."],
        ["D. Konsistensi", "Perjalanan = bisnis + pribadi; jumlah semua mitra = dunia", "Penyesuaian kecil."],
        ["E. Penyeimbangan", "Nilai seimbang (B) = rata-rata tertimbang estimasi sisi Indonesia dan laporan mitra",
         "Bobot = indeks simetri pelapor (1 − asimetri, rata-rata 3 tahun) × bobot keyakinan metode (tabel bawah)."],
        ["F. Ekstrapolasi", "Tahun terbaru diekstrapolasi dari deret waktu",
         "2023–2024: sebagian besar baris Indonesia berkode IE."],
    ]
    end = write_rows(ws, top + 2, 1, steps, [None, None, None])
    for rr in range(top + 2, end):
        for cc in (2, 3):
            ws.cell(row=rr, column=cc).alignment = Alignment(wrap_text=True, vertical="top")
        ws.row_dimensions[rr].height = 45
    r = end + 1
    bold(ws, r, 1, "Bobot keyakinan metode (GIVEN, dari dokumen OECD-WTO)", 10)
    header(ws, r + 1, 1, ["Jenis data", "Bobot"], height=20)
    r = write_rows(ws, r + 2, 1, [["Dilaporkan", 1.0], ["Deret waktu (interpolasi/ekstrapolasi)", 0.75],
                                  ["Model gravitasi", 0.5], ["Dilaporkan tapi tidak masuk akal", 0.25]],
                   [None, "0.00"])
    r += 1
    r = code_table(ws, r, [list(x) for x in BATIS_CODES], "Arti kode di data (label resmi OECD-WTO)",
                   spans=((1, 1), (2, 2), (3, 3), (4, 4)))
    ws.column_dimensions["D"].width = 60
    b23 = batis_year(batis, 2023)
    ex = b23.loc["AUS"]
    w = (ex.kredit_B - ex.kredit_F) / (ex.kredit_mitra - ex.kredit_F) if pd.notna(ex.kredit_mitra) else None
    r += 1
    bold(ws, r, 1, "Contoh dari data: kredit Indonesia dari Australia, 2023 (USD juta)", 10)
    first = header(ws, r + 1, 1, ["Komponen", "Nilai", "Status"], height=20)
    r = write_rows(ws, first, 1, [["Estimasi OECD-WTO sisi Indonesia (F)", ex.kredit_F, "GIVEN (OECD-WTO)"],
                                  ["Laporan Australia (N)", ex.kredit_mitra, "GIVEN (OECD-WTO; asal ABS)"],
                                  ["Nilai seimbang (B)", ex.kredit_B, "GIVEN (OECD-WTO)"],
                                  ["Bobot tersirat laporan Australia", w,
                                   "SELF-CALCULATED: (B − F) ÷ (N − F), hanya untuk ilustrasi"]],
                   [None, NUM, None])
    ws.cell(row=r - 1, column=2).number_format = "0%"
    r += 1
    bold(ws, r, 1, "Batas model", 10)
    c = ws.cell(row=r + 1, column=1,
                value=("Model gravitasi tidak tahu soal haji dan umrah. Debit Indonesia ke Arab Saudi menurut BaTIS "
                       f"hanya USD {bt.loc['SAU', 'debit_B']:,.0f} juta ({YEAR}); hitungan BPS sekitar USD "
                       f"{bps.loc['SAU', 'debit']:,.0f} juta dan biaya haji saja (di luar tiket) sekitar USD "
                       f"{ds.loc['SAU', 'debit_usd']:,.0f} juta. Karena itu BaTIS hanya dipakai sebagai pembanding."))
    c.font = Font(name=FONT, size=10)
    ws.merge_cells(start_row=r + 1, start_column=1, end_row=r + 1, end_column=3)
    c.alignment = Alignment(wrap_text=True, vertical="top")
    ws.row_dimensions[r + 1].height = 32

    # ================================================================ 9. Sisi Tujuan
    ws = wb.create_sheet("Data Negara Tujuan")
    style_ws(ws)
    title(ws, f"Pembanding: penerimaan negara tujuan dari wisatawan Indonesia, {YEAR}",
          "Penerimaan negara tujuan dari wisatawan Indonesia = debit Indonesia. Hanya satu sisi neraca. "
          "Dikonversi ke USD dengan kurs Bank Dunia.")
    prov_legend(ws)
    d = ds.sort_values("debit_usd", ascending=False)
    status = {
        "Tourism Malaysia": f"Nilai lokal GIVEN. USD SELF-CALCULATED: RM × (1 − 2,2% tiket internasional) ÷ "
                            f"{FX[('MYS', YEAR)]:.4f} (kurs MYR per USD)",
        "STB": f"Nilai lokal GIVEN. USD SELF-CALCULATED: S$ ÷ {FX[('SGP', YEAR)]:.4f} (kurs SGD per USD)",
        "JTA": f"GIVEN ¥1.109 oku (= ¥110.900 juta). USD SELF-CALCULATED: ÷ {FX[('JPN', YEAR)]:.2f} (kurs JPY per USD)",
        "MOTS": "GIVEN: nilai USD langsung dari publikasi MOTS (tidak dikonversi ulang)",
        "Haji": f"SELF-CALCULATED: kuota haji {YEAR} (GIVEN) × (BPIH rata-rata − biaya tiket) ÷ {FX[('IDN', YEAR)]:,.0f}"
                f" (kurs IDR). SELF-ASSUMED: biaya tiket final 2025 (Rp33,1 juta) dipakai untuk {YEAR}.",
    }
    first = header(ws, 4, 1, ["Negara tujuan", "Nilai (mata uang lokal)", "Satuan", "Debit Indonesia (USD juta)",
                              "Periode", "Status angka", "Catatan", "Halaman data", "Unduh langsung"],
                   [16, 14, 10, 14, 12, 52, 40, 22, 24],
                   height=30)
    end = write_rows(ws, first, 1, [[r_.nama, r_.nilai_lokal, r_.mata_uang, r_.debit_usd, r_.periode,
                                     status[r_.sumber], r_.catatan, None, None] for _, r_ in d.iterrows()],
                     [None, NUM, None, NUM, None, None, None, None, None])
    for i, (_, r_) in enumerate(d.iterrows()):
        s_ = SRC[r_.sumber]
        link(ws.cell(row=first + i, column=8), s_["browse"][0][1], s_["browse"][0][0])
        link(ws.cell(row=first + i, column=9), s_["unduh"][0][1], s_["unduh"][0][0])
        ws.cell(row=first + i, column=6).alignment = Alignment(wrap_text=True, vertical="top")
        ws.row_dimensions[first + i].height = 40
    src_end = src_block(ws, end + 2, list(dict.fromkeys(d.sumber)) + ["WB_KURS"])
    add_charts(ws, pd.DataFrame({"nama": d.nama, "debit": d.debit_usd}), 1, src_end + 1, "Sumber: negara tujuan",
               flows=("debit",))

    # ================================================================ 10. Tren
    ws = wb.create_sheet("Tren")
    style_ws(ws)
    title(ws, "Tren neraca Indonesia per negara (USD juta)",
          "BPS hanya tersedia untuk tahun publikasi; untuk deret panjang dipakai dua pembanding: mirror OECD TIS "
          "(negara pelapor) dan BaTIS (semua negara, hasil model).")
    years = list(range(2010, YEAR + 1))
    tis_all = tis.pivot_table(index=["reporter", "year"], columns="flow_idn", values="value", aggfunc="sum")
    tis_bal = (tis_all["credit"] - tis_all["debit"]).unstack("year")
    prov_legend(ws)
    bold(ws, 4, 1, "Mirror OECD TIS (10 negara dengan neraca terbesar)", 10)
    first = header(ws, 5, 1, ["Negara"] + years, [20] + [9] * len(years), height=20,
                   prov="SELF-CALCULATED – MIRRORING: kredit − debit Indonesia; kredit dan debit GIVEN – MIRRORING dari OECD TIS",
                   prov_height=18)
    top_rep = tis_y.neraca.abs().sort_values(ascending=False).head(10).index
    r = write_rows(ws, first, 1, [[nama(i)] + [tis_bal.loc[i, y] if y in tis_bal.columns else None for y in years]
                                  for i in top_rep], [None] + [NUM0] * len(years))
    bt_hist = batis[(batis.line == "A_model") & batis.adjustment.str.startswith("B") & (batis.partner_iso3 != "W")]
    bh = bt_hist.pivot_table(index=["partner_iso3", "year"], columns="flow_idn", values="value", aggfunc="sum")
    bh = (bh["credit"] - bh["debit"]).unstack("year")
    top_bt = bt_std.neraca.abs().sort_values(ascending=False).head(10).index
    r += 1
    bold(ws, r, 1, "BaTIS nilai seimbang (10 negara dengan neraca terbesar)", 10)
    first = header(ws, r + 1, 1, ["Negara"] + years, height=20,
                   prov="Kredit B dan debit B GIVEN (diambil langsung dari OECD-WTO BaTIS, kode ME/IE: lihat sheet Legenda); neraca = SELF-CALCULATED: kredit B − debit B saja",
                   prov_height=18)
    write_rows(ws, first, 1, [[bt_std.loc[i, "nama"]] + [bh.loc[i, y] if y in bh.columns else None for y in years]
                              for i in top_bt], [None] + [NUM0] * len(years))

    # ================================================================ 11. Rekonsiliasi
    ws = wb.create_sheet("Rekonsiliasi")
    style_ws(ws)
    title(ws, "Dibandingkan dengan total resmi Bank Indonesia (USD juta)",
          "BI SEKI Tabel V.3 = angka resmi jasa perjalanan Indonesia dengan seluruh dunia (tanpa rincian negara).")
    prov_legend(ws)
    first = header(ws, 4, 1, ["Sisi", f"BPS {YEAR}", f"BI {YEAR}", "BPS / BI"], [30, 14, 14, 12], height=20,
                   prov=[None, "SELF-CALCULATED: total di sheet BPS – Perhitungan", "GIVEN (BI SEKI Tabel V.3)",
                         "SELF-CALCULATED: BPS ÷ BI"])
    r = write_rows(ws, first, 1, [["Kredit", bc, bi.loc[YEAR, "credit"], bc / bi.loc[YEAR, "credit"]],
                              ["Debit", bd, bi.loc[YEAR, "debit"], bd / bi.loc[YEAR, "debit"]],
                              ["Neraca", bc - bd, bi.loc[YEAR, "balance"], None]], [None, NUM, NUM, PCT])
    note(ws, r, "BPS lebih tinggi dari BI; BI tidak menerbitkan rincian penyesuaiannya. Porsi dan urutan negara dari "
                "BPS tetap bisa dipakai.")
    bsum = bt_hist.pivot_table(index="year", columns="flow_idn", values="value", aggfunc="sum")
    eu_only = eu[~eu.reporter.isin(tis.reporter.unique())]
    mir = pd.concat([tis[["year", "flow_idn", "value"]],
                     eu_only[["year", "flow_idn", "value_usd"]].rename(columns={"value_usd": "value"})])
    msum = mir[mir.flow_idn.isin(["credit", "debit"])].pivot_table(index="year", columns="flow_idn", values="value",
                                                                    aggfunc="sum")
    r += 2
    first = header(ws, r, 1, ["Tahun", "BI kredit", "BI debit", "BI neraca", "BaTIS kredit (jumlah semua mitra)",
                      "BaTIS debit (jumlah semua mitra)", "Mirror kredit (jumlah negara pelapor)", "Porsi dari BI",
                      "Mirror debit (jumlah negara pelapor)", "Porsi dari BI"], [30, 11, 11, 11, 14, 14, 15, 9, 15, 9],
                   prov=[None, "GIVEN (BI)", "GIVEN (BI)", "GIVEN (BI)",
                         "SELF-CALCULATED: jumlah semua mitra (GIVEN BaTIS B)", "SELF-CALCULATED: jumlah semua mitra",
                         "SELF-CALCULATED: jumlah pelapor OECD TIS + Eurostat-saja (EUR→USD)", "SELF-CALCULATED: ÷ BI",
                         "SELF-CALCULATED: jumlah pelapor", "SELF-CALCULATED: ÷ BI"], prov_height=64)
    rows = [[y, bi.loc[y, "credit"], bi.loc[y, "debit"], bi.loc[y, "balance"], bsum.loc[y, "credit"],
             bsum.loc[y, "debit"], msum.loc[y, "credit"], msum.loc[y, "credit"] / bi.loc[y, "credit"],
             msum.loc[y, "debit"], msum.loc[y, "debit"] / bi.loc[y, "debit"]] for y in range(2018, YEAR + 1)]
    r = write_rows(ws, first, 1, rows, [None, NUM, NUM, NUM, NUM, NUM, NUM, PCT, NUM, PCT])
    note(ws, r, "Mirror hanya mencakup negara pelapor, jadi porsinya di bawah 100%.")
    src_block(ws, r + 2, ["BI", "BATIS", "OECD_TIS", "EUROSTAT"])

    # ================================================================ 12. Sumber
    ws = wb.create_sheet("Sumber")
    style_ws(ws)
    title(ws, "Daftar sumber yang dicek",
          "Halaman data = halaman untuk browse data / menu unduh. Unduh langsung = klik langsung mengunduh file. "
          "Kolom 'Cek link' = hasil membuka setiap link (detail di sheet 'Cek Link'). Bukti lengkap: folder 01_discovery.")
    chk = link_status()
    inv = pd.read_csv(ROOT / "01_discovery" / "source_inventory.csv")
    kelas = [("FULL BALANCE (MODELLED)", "Kredit + debit (hasil model)"), ("FULL BALANCE", "Kredit + debit"),
             ("ONE SIDE", "Satu sisi saja"), ("COUNT ONLY", "Jumlah orang (bukan nilai)"),
             ("TOTAL ONLY", "Total tanpa rincian negara"), ("NOT AVAILABLE", "Tidak ada rincian negara"),
             ("LOCKED", "Perlu login"), ("NARRATIVE", "Narasi saja"), ("DERIVED", "Turunan (duplikat)"),
             ("TO CHECK", "Belum dicek")]

    def k(x):
        return next((b_ for a_, b_ in kelas if str(x).startswith(a_)), x)
    header(ws, 4, 1, ["Sumber", "Penerbit", "Isi data", "Status", "Halaman data", "Unduh langsung", "Cek link"],
           [48, 30, 26, 44, 30, 30, 34], height=20)

    def cek(urls):
        v = [chk.get(u) for u in urls]
        if not v or any(x is None for x in v):
            return "Belum dicek"
        bad = [x[0] for x in v if not x[0].startswith("OK")]
        if bad:
            return "; ".join(bad)
        anti = any("anti-bot" in x[0] for x in v)
        return f"OK ({v[0][1]})" + ("; halaman BPS: buka di browser biasa" if anti else "")
    r = 5
    inv_rows = [(r_.source, r_.publisher, k(r_.balance_class), r_.status, INVENTORY[r_.source_id])
                for _, r_ in inv.iterrows()] + [("Kurs Bank Dunia (rata-rata tahunan)", "World Bank", "Kurs", "Dipakai",
                                                 "WB_KURS")]
    for src_, pub_, isi_, st_, key in inv_rows:
        s_ = SRC[key]
        b_, u_ = s_["browse"], s_["unduh"]
        write_rows(ws, r, 1, [[src_, pub_, isi_, st_, None, None, cek([x[1] for x in b_ + u_])]], [None] * 7)
        for j in range(max(len(b_), len(u_), 1)):
            if j < len(b_):
                link(ws.cell(row=r + j, column=5), b_[j][1], b_[j][0])
            if j < len(u_):
                link(ws.cell(row=r + j, column=6), u_[j][1], u_[j][0])
            elif j == 0:
                note(ws, r, "– (tidak ada file unduhan langsung)", col=6)
        if s_.get("catatan"):
            note(ws, r + max(len(b_), len(u_), 1), s_["catatan"], col=5)
            r += 1
        r += max(len(b_), len(u_), 1)

    # ================================================================ 13. Cek Link
    chk = link_status()
    seen = {}
    for w in wb.worksheets:
        for row_ in w.iter_rows():
            for c in row_:
                if c.hyperlink and c.hyperlink.target:
                    seen.setdefault(c.hyperlink.target, []).append(f"{w.title}!{c.coordinate}")
    ws = wb.create_sheet("Cek Link")
    style_ws(ws)
    title(ws, "Hasil cek semua link di workbook ini",
          "Setiap link dibuka otomatis (scripts/check_links.py): permintaan biasa dulu, lalu browser Chromium. "
          "bps.go.id memakai verifikasi anti-bot (Cloudflare) sehingga tidak bisa diuji otomatis; di browser biasa terbuka "
          "normal. Link unduhan BPS (webapi.bps.go.id) bisa diuji dan berhasil.")
    first = header(ws, 4, 1, ["Link", "Hasil cek", "Tanggal cek", "Dipakai di"], [70, 40, 12, 40], height=20)
    r = first
    for url in sorted(seen):
        v, dte = chk.get(url, ("Belum dicek", None))
        write_rows(ws, r, 1, [[None, v, dte, ", ".join(seen[url][:4]) + (" …" if len(seen[url]) > 4 else "")]],
                   [None] * 4)
        link(ws.cell(row=r, column=1), url, url if len(url) < 110 else url[:107] + "…")
        if not v.startswith("OK"):
            ws.cell(row=r, column=2).font = Font(name=FONT, size=10, bold=True, color="C00000")
        r += 1

    wb.calculation.fullCalcOnLoad = True
    OUT.parent.mkdir(exist_ok=True)
    wb.save(OUT)

    # ---------- QA output
    lines = [f"# Reconciliation {YEAR} (build_excel.py)", "",
             f"BI {YEAR}: credit {bi.loc[YEAR, 'credit']:,.1f}, debit {bi.loc[YEAR, 'debit']:,.1f}",
             f"BPS {YEAR}: credit {bc:,.1f} ({bc / bi.loc[YEAR, 'credit']:.0%} of BI), debit {bd:,.1f} "
             f"({bd / bi.loc[YEAR, 'debit']:.0%} of BI)", "",
             "Consensus top-5 (>=3 of 5 sources): " + str({k_: sorted(v) for k_, v in consensus.items()})]
    for name, dfx in sources.items():
        for flow in ("kredit", "debit", "neraca"):
            t = top5(dfx, flow)
            if len(t):
                lines.append(f"- {name} {flow}: " + ", ".join(f"{dfx.loc[i, 'nama']} {v:,.0f}" for i, v in t.items()))
    (ROOT / "04_qa" / "reconciliation.md").write_text("\n".join(lines), encoding="utf-8")
    print("\n".join(lines))
    print("saved", OUT)


if __name__ == "__main__":
    main()
